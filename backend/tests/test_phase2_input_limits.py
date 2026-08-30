import io
import uuid

import pytest
from docx import Document as DocxDocument
from pydantic import ValidationError

from app.auth import dependencies as auth_deps
from app.database.connection import AsyncSessionLocal
from app.models.document import Document
from app.models.session import ChatSession
from app.models.user import User
from app.schemas.chat import ChatRequest
from app.schemas.search import SearchRequest
from ai.services.parser_service import parse_document


# --- 1. Chat message / image_data size limits ---


def test_chat_message_within_limit_is_accepted():
    req = ChatRequest(session_id=uuid.uuid4(), message="A reasonable question.")
    assert req.message == "A reasonable question."


def test_chat_message_over_limit_is_rejected():
    with pytest.raises(ValidationError):
        ChatRequest(session_id=uuid.uuid4(), message="x" * 8001)


def test_chat_image_data_over_limit_is_rejected():
    with pytest.raises(ValidationError):
        ChatRequest(
            session_id=uuid.uuid4(),
            message="describe this",
            image_data="x" * 20_000_000,
        )


# --- 2. Search query max length ---


def test_search_query_within_limit_is_accepted():
    req = SearchRequest(session_id=uuid.uuid4(), query="normal query")
    assert req.query == "normal query"


def test_search_query_over_limit_is_rejected():
    with pytest.raises(ValidationError):
        SearchRequest(session_id=uuid.uuid4(), query="x" * 1001)


# --- 3. Bounded _last_sync_time cache ---


def test_last_sync_cache_evicts_oldest_beyond_max_size():
    original_cache = auth_deps._last_sync_time
    original_max = auth_deps._MAX_SYNC_CACHE_ENTRIES
    try:
        auth_deps._last_sync_time = type(original_cache)()
        auth_deps._MAX_SYNC_CACHE_ENTRIES = 5

        for i in range(10):
            auth_deps._record_sync(f"user-{i}", float(i))

        assert len(auth_deps._last_sync_time) == 5
        # The 5 most-recently-synced users survive; earliest ones are evicted.
        assert "user-0" not in auth_deps._last_sync_time
        assert "user-4" not in auth_deps._last_sync_time
        assert "user-9" in auth_deps._last_sync_time
        assert "user-5" in auth_deps._last_sync_time
    finally:
        auth_deps._last_sync_time = original_cache
        auth_deps._MAX_SYNC_CACHE_ENTRIES = original_max


# --- 4. .doc upload fails fast with a clear message; .docx still works ---


@pytest.mark.asyncio
async def test_doc_upload_fails_fast_with_clear_message():
    with pytest.raises(ValueError, match=r"\.doc.*not supported"):
        await parse_document(b"not a real ole2 doc file", "legacy.doc")


@pytest.mark.asyncio
async def test_docx_upload_still_parses_correctly():
    buf = io.BytesIO()
    docx_doc = DocxDocument()
    docx_doc.add_paragraph("Hello from a real docx test fixture.")
    docx_doc.save(buf)

    pages = await parse_document(buf.getvalue(), "real.docx")
    assert len(pages) == 1
    assert "Hello from a real docx test fixture." in pages[0]["text"]


# --- 5. Document/DocumentChunk created_at is timezone-aware ---


@pytest.mark.asyncio
async def test_document_created_at_is_timezone_aware():
    async with AsyncSessionLocal() as db:
        user = User(id=uuid.uuid4(), email=f"phase2-{uuid.uuid4()}@example.com")
        db.add(user)
        await db.flush()

        session = ChatSession(id=uuid.uuid4(), user_id=user.id, title="Phase2 tz test")
        db.add(session)
        await db.flush()

        doc = Document(
            id=uuid.uuid4(),
            user_id=user.id,
            session_id=session.id,
            filename="test.txt",
            file_type="txt",
            file_path="test/test.txt",
            file_size=10,
            status="processing",
        )
        db.add(doc)
        await db.commit()
        await db.refresh(doc)

        assert doc.created_at.tzinfo is not None

        # ON DELETE CASCADE on the FKs cleans up the session and document.
        await db.delete(user)
        await db.commit()
