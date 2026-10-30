import io
import uuid
import pytest
from unittest.mock import AsyncMock, MagicMock, patch
from docx import Document

from ai.services.chunking_service import chunk_text
from ai.services.parser_service import (
    extract_docx_headers_and_footers,
    format_docx_table,
    iter_docx_body_blocks,
    parse_document,
)
from app.models.document import Document as DBDocument
from app.routes.upload import process_document_background


def create_docx_bytes(build_fn) -> bytes:
    """Helper to create in-memory DOCX file bytes from a builder function."""
    doc = Document()
    build_fn(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ===========================================================================
# GROUP 1: Paragraph and Table Order Preservation
# ===========================================================================

@pytest.mark.asyncio
async def test_docx_paragraphs_and_table_order_preserved():
    """
    Verify that paragraphs and tables are extracted in their exact document order.
    Issue #19: Tables should not be dropped, and order must be maintained.
    """
    def build(doc):
        doc.add_paragraph("First paragraph: Introduction.")
        tbl1 = doc.add_table(rows=2, cols=2)
        tbl1.cell(0, 0).text = "Metric"
        tbl1.cell(0, 1).text = "Value"
        tbl1.cell(1, 0).text = "Active Users"
        tbl1.cell(1, 1).text = "10,000"
        doc.add_paragraph("Second paragraph: Transition between tables.")
        tbl2 = doc.add_table(rows=2, cols=2)
        tbl2.cell(0, 0).text = "Department"
        tbl2.cell(0, 1).text = "Headcount"
        tbl2.cell(1, 0).text = "Engineering"
        tbl2.cell(1, 1).text = "45"
        doc.add_paragraph("Third paragraph: Concluding remarks.")

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "report.docx")

    assert len(pages) == 1
    text = pages[0]["text"]

    # Verify all expected content is present
    assert "First paragraph: Introduction." in text
    assert "| Metric | Value |" in text
    assert "| Active Users | 10,000 |" in text
    assert "Second paragraph: Transition between tables." in text
    assert "| Department | Headcount |" in text
    assert "| Engineering | 45 |" in text
    assert "Third paragraph: Concluding remarks." in text

    # Strict sequential ordering verification
    pos_p1 = text.find("First paragraph: Introduction.")
    pos_t1 = text.find("| Metric | Value |")
    pos_p2 = text.find("Second paragraph: Transition between tables.")
    pos_t2 = text.find("| Department | Headcount |")
    pos_p3 = text.find("Third paragraph: Concluding remarks.")

    assert pos_p1 < pos_t1 < pos_p2 < pos_t2 < pos_p3


# ===========================================================================
# GROUP 2: Table Formatting and Edge Cases
# ===========================================================================

@pytest.mark.asyncio
async def test_docx_table_markdown_formatting():
    """
    Verify that DOCX tables are converted into valid Markdown table syntax,
    including header row, separator row, and data rows.
    """
    def build(doc):
        tbl = doc.add_table(rows=3, cols=3)
        headers = ["Product", "Price", "Stock"]
        row1 = ["Laptop", "$1,200", "50"]
        row2 = ["Mouse", "$25", "200"]

        for col_idx, h in enumerate(headers):
            tbl.cell(0, col_idx).text = h
        for col_idx, v in enumerate(row1):
            tbl.cell(1, col_idx).text = v
        for col_idx, v in enumerate(row2):
            tbl.cell(2, col_idx).text = v

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "pricing.docx")
    text = pages[0]["text"]

    expected_table = (
        "| Product | Price | Stock |\n"
        "| --- | --- | --- |\n"
        "| Laptop | $1,200 | 50 |\n"
        "| Mouse | $25 | 200 |"
    )
    assert expected_table in text


@pytest.mark.asyncio
async def test_docx_table_empty_cells_handled():
    """
    Verify that empty cells in tables are preserved as empty cells in markdown
    without crashing or causing column misalignment.
    """
    def build(doc):
        tbl = doc.add_table(rows=2, cols=3)
        tbl.cell(0, 0).text = "Col A"
        tbl.cell(0, 1).text = "Col B"
        tbl.cell(0, 2).text = "Col C"

        tbl.cell(1, 0).text = "Val A"
        # cell(1, 1) is intentionally empty
        tbl.cell(1, 2).text = "Val C"

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "sparse_table.docx")
    text = pages[0]["text"]

    expected_row = "| Val A |  | Val C |"
    assert expected_row in text


@pytest.mark.asyncio
async def test_docx_table_cell_newlines_and_pipes_escaped():
    """
    Verify that multiline cells have newlines replaced by spaces
    and literal pipe characters are escaped to preserve table row formatting.
    """
    def build(doc):
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "Name"
        tbl.cell(0, 1).text = "Description"

        tbl.cell(1, 0).text = "Widget | Pro"
        tbl.cell(1, 1).text = "Line 1\nLine 2\nLine 3"

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "multiline.docx")
    text = pages[0]["text"]

    assert "Widget \\| Pro" in text
    assert "Line 1 Line 2 Line 3" in text


@pytest.mark.asyncio
async def test_docx_horizontally_merged_cells():
    """
    Verify that horizontally merged cells (which share the same _tc element)
    do not duplicate text across columns and preserve total column count.
    """
    def build(doc):
        tbl = doc.add_table(rows=2, cols=3)
        tbl.cell(0, 0).merge(tbl.cell(0, 1))
        tbl.cell(0, 0).text = "Merged Header"
        tbl.cell(0, 2).text = "Col 3"

        tbl.cell(1, 0).text = "Val 1"
        tbl.cell(1, 1).text = "Val 2"
        tbl.cell(1, 2).text = "Val 3"

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "merged.docx")
    text = pages[0]["text"]

    # Header row should have "Merged Header" followed by an empty cell for the merged span
    assert "| Merged Header |  | Col 3 |" in text
    assert "| Val 1 | Val 2 | Val 3 |" in text


@pytest.mark.asyncio
async def test_docx_empty_table_ignored():
    """
    Verify that tables with no rows or only whitespace cells are skipped
    and do not output empty Markdown artifacts.
    """
    def build(doc):
        doc.add_paragraph("Paragraph before.")
        doc.add_table(rows=2, cols=2)  # All cells empty
        doc.add_paragraph("Paragraph after.")

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "empty_tbl.docx")
    text = pages[0]["text"]

    assert text == "Paragraph before.\n\nParagraph after."


# ===========================================================================
# GROUP 3: Headers and Footers Extraction & Deduplication
# ===========================================================================

@pytest.mark.asyncio
async def test_docx_headers_and_footers_extracted():
    """
    Verify that document headers appear at the top and footers at the bottom.
    """
    def build(doc):
        sec = doc.sections[0]
        sec.header.paragraphs[0].text = "Header: Financial Quarterly Report"
        sec.footer.paragraphs[0].text = "Footer: Confidential & Proprietary"
        doc.add_paragraph("Body paragraph content.")

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "doc_with_hf.docx")
    text = pages[0]["text"]

    pos_header = text.find("Header: Financial Quarterly Report")
    pos_body = text.find("Body paragraph content.")
    pos_footer = text.find("Footer: Confidential & Proprietary")

    assert pos_header != -1
    assert pos_body != -1
    assert pos_footer != -1
    assert pos_header < pos_body < pos_footer


@pytest.mark.asyncio
async def test_docx_headers_and_footers_deduplication():
    """
    Verify that multi-section documents with linked or repeated headers/footers
    do not duplicate header/footer text blocks.
    """
    def build(doc):
        sec1 = doc.sections[0]
        sec1.header.paragraphs[0].text = "Company Confidential"
        sec1.footer.paragraphs[0].text = "Page Footer 2026"
        doc.add_paragraph("Section 1 Body")

        sec2 = doc.add_section()
        # sec2 header/footer is linked to previous by default
        doc.add_paragraph("Section 2 Body")

        sec3 = doc.add_section()
        # Unlinked header with same text
        sec3.header.is_linked_to_previous = False
        sec3.header.paragraphs[0].text = "Company Confidential"
        doc.add_paragraph("Section 3 Body")

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "multisection.docx")
    text = pages[0]["text"]

    # "Company Confidential" should only appear ONCE in headers
    assert text.count("Company Confidential") == 1
    # "Page Footer 2026" should only appear ONCE in footers
    assert text.count("Page Footer 2026") == 1


@pytest.mark.asyncio
async def test_docx_first_page_header_footer():
    """
    Verify extraction when different first page header/footer is enabled.
    """
    def build(doc):
        sec = doc.sections[0]
        sec.different_first_page_header_footer = True
        sec.first_page_header.paragraphs[0].text = "Cover Page Header"
        sec.first_page_footer.paragraphs[0].text = "Cover Page Footer"
        sec.header.paragraphs[0].text = "Main Body Header"
        sec.footer.paragraphs[0].text = "Main Body Footer"
        doc.add_paragraph("Body text.")

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "first_page_hf.docx")
    text = pages[0]["text"]

    assert "Cover Page Header" in text
    assert "Main Body Header" in text
    assert "Cover Page Footer" in text
    assert "Main Body Footer" in text


# ===========================================================================
# GROUP 4: Regression and Edge Cases
# ===========================================================================

@pytest.mark.asyncio
async def test_docx_text_only_backward_compatibility():
    """
    Verify that existing DOCX files containing only text paragraphs
    continue to parse identically to the previous implementation.
    """
    def build(doc):
        doc.add_paragraph("Paragraph 1.")
        doc.add_paragraph("Paragraph 2.")
        doc.add_paragraph("Paragraph 3.")

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "plain.docx")

    assert len(pages) == 1
    assert pages[0]["text"] == "Paragraph 1.\n\nParagraph 2.\n\nParagraph 3."


@pytest.mark.asyncio
async def test_docx_empty_file_and_empty_document():
    """Verify that empty bytes and empty documents return clean empty text without crashing."""
    empty_bytes_pages = await parse_document(b"", "empty.docx")
    assert empty_bytes_pages == [{"page_number": 1, "text": ""}]

    blank_doc_bytes = create_docx_bytes(lambda doc: None)
    blank_doc_pages = await parse_document(blank_doc_bytes, "blank.docx")
    assert blank_doc_pages == [{"page_number": 1, "text": ""}]


# ===========================================================================
# GROUP 5: Chunking and Ingestion Pipeline Integration
# ===========================================================================

@pytest.mark.asyncio
async def test_docx_table_chunks_correctly():
    """
    Verify that extracted DOCX table text reaches the chunking service intact
    and semantic boundaries preserve table readability.
    """
    def build(doc):
        doc.add_paragraph("System Performance Overview:")
        tbl = doc.add_table(rows=3, cols=3)
        tbl.cell(0, 0).text = "Service"
        tbl.cell(0, 1).text = "Latency"
        tbl.cell(0, 2).text = "Status"

        tbl.cell(1, 0).text = "API Gateway"
        tbl.cell(1, 1).text = "12ms"
        tbl.cell(1, 2).text = "Healthy"

        tbl.cell(2, 0).text = "Vector DB"
        tbl.cell(2, 1).text = "45ms"
        tbl.cell(2, 2).text = "Healthy"

        doc.add_paragraph("All services are operating within acceptable parameters.")

    docx_bytes = create_docx_bytes(build)
    pages = await parse_document(docx_bytes, "system.docx")
    text = pages[0]["text"]

    chunks = chunk_text(text, chunk_size=600, chunk_overlap=120)
    assert len(chunks) >= 1
    first_chunk_content = chunks[0]["content"]

    assert "| Service | Latency | Status |" in first_chunk_content
    assert "| API Gateway | 12ms | Healthy |" in first_chunk_content
    assert "| Vector DB | 45ms | Healthy |" in first_chunk_content


@pytest.mark.asyncio
async def test_docx_end_to_end_ingestion_with_tables():
    """
    End-to-end integration test verifying that a DOCX with tables and paragraphs
    progresses through document parsing, chunking, embedding generation,
    and database persistence with status 'ready'.
    """
    doc_id = uuid.uuid4()

    def build(doc):
        doc.add_paragraph("Employee Directory:")
        tbl = doc.add_table(rows=2, cols=3)
        tbl.cell(0, 0).text = "Name"
        tbl.cell(0, 1).text = "Role"
        tbl.cell(0, 2).text = "Location"

        tbl.cell(1, 0).text = "Alice Chen"
        tbl.cell(1, 1).text = "Principal Architect"
        tbl.cell(1, 2).text = "San Francisco"

    docx_bytes = create_docx_bytes(build)

    mock_doc = MagicMock(spec=DBDocument)
    mock_doc.id = doc_id
    mock_doc.status = "processing"
    mock_doc.ocr_warning_count = 0
    mock_doc.error_message = None

    async def mock_execute(stmt):
        mock_res = MagicMock()
        stmt_str = str(stmt).lower()
        if "count(" in stmt_str:
            mock_res.scalar.return_value = 1
        else:
            mock_res.scalar_one_or_none.return_value = mock_doc
        return mock_res

    mock_session = AsyncMock()
    mock_session.add = MagicMock()
    mock_session.execute.side_effect = mock_execute

    with patch("app.routes.upload.AsyncSessionLocal") as mock_session_maker, \
         patch("app.routes.upload.embedding_service") as mock_es:

        mock_session_maker.return_value.__aenter__.return_value = mock_session
        mock_es.generate_batch_embeddings = AsyncMock(return_value=[[0.1] * 1536])

        await process_document_background(doc_id, docx_bytes, "directory.docx", "path/directory.docx")

        # Verify chunk was persisted with table contents
        add_calls = mock_session.add.call_args_list
        assert len(add_calls) >= 1
        saved_chunk = add_calls[0][0][0]
        assert "Alice Chen" in saved_chunk.content
        assert "Principal Architect" in saved_chunk.content
        assert "| Name | Role | Location |" in saved_chunk.content

        # Verify document transitioned to ready
        assert mock_doc.status == "ready"
        assert mock_doc.error_message is None
        mock_session.commit.assert_called()
